"""A deck built from code or from the [[slide]] tables of a TOML file.

    [[slide]]
    title = "Bessel functions"
    [[slide.item]]
    text = "Generating function"
    [[slide.item]]
    math = 'e^{\\frac{x}{2}(t-1/t)} = \\sum_n J_n(x)\\,t^n'
    route = "tex"             # this equation by the other route than [math] route
    [[slide.item]]
    figure = "bessel.svg"     # PNG fallback: bessel.png next to it, or fallback = "..."
    width = 6                 # inches; any item takes left, top, width

Items without a position flow down the slide below the title.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from . import omml, tex
from .config import Config
from .picture import add_svg, fit

MARGIN = 0.6
GAP = 0.15
# Average glyph width of a proportional sans font, in em, to guess where text wraps.
CHAR_EM = 0.5
SLIDE_KEYS = {"title", "item"}
ITEM_KEYS = {"text": {"size", "bold"}, "math": {"route", "size"}, "figure": {"fallback", "height"}}
PLACE_KEYS = {"left", "top", "width"}


def add_equation(
    slide,
    latex: str,
    *,
    left,
    top,
    route: str = "omml",
    font: str | None = None,
    size_pt: float = 24,
    width=None,
):
    """Add `latex` to `slide` at (left, top) and return the shape. route "omml" makes an editable
    Office equation, "tex" an SVG picture through xelatex. `width` caps the shape's width."""
    if route == "omml":
        return omml.add_equation(
            slide,
            latex,
            font=font or "Cambria Math",
            size_pt=size_pt,
            left=left,
            top=top,
            width=width or Inches(8),
        )
    if route == "tex":
        return tex.add_equation(
            slide, latex, font=font or "STIX Two Math", size_pt=size_pt, left=left, top=top, width=width
        )
    raise ValueError(f"route must be 'omml' or 'tex', not {route!r}")


def add_figure(
    slide, path: str | Path, *, left, top, width=None, height=None, fallback: str | Path | None = None
):
    """A PNG or JPEG picture, or an SVG picture with a PNG fallback (`fallback`, else the .png
    next to the SVG)."""
    path = Path(path)
    if path.suffix.lower() != ".svg":
        return slide.shapes.add_picture(str(path), left, top, width, height)
    png = Path(fallback) if fallback else path.with_suffix(".png")
    if not png.exists():
        raise FileNotFoundError(f"{path.name} needs a PNG fallback: {png.name} not found")
    return add_svg(slide, path.read_bytes(), png.read_bytes(), left=left, top=top, width=width, height=height)


class Deck:
    """A presentation with the configured size, fonts and math route."""

    def __init__(self, config: Config | None = None):
        self.config = config or Config()
        template = self.config.deck.template
        self.prs = Presentation(str(self.config.root / template) if template else None)
        if not template:
            self.prs.slide_width = Inches(self.config.deck.width)
            self.prs.slide_height = Inches(self.config.deck.height)
        self._cursor: dict[int, float] = {}

    @property
    def width(self) -> float:
        return self.prs.slide_width / 914400

    def slide(self, title: str = ""):
        """A blank slide, with `title` at the top if given."""
        layout = self.prs.slide_layouts[6 if len(self.prs.slide_layouts) > 6 else -1]
        slide = self.prs.slides.add_slide(layout)
        self._cursor[slide.slide_id] = MARGIN
        if title:
            self.text(slide, title, size=self.config.deck.title_size, bold=True)
        return slide

    def _place(self, slide, left, top, width) -> tuple[float, float, float]:
        y = self._cursor.get(slide.slide_id, MARGIN)
        return (
            MARGIN if left is None else left,
            y if top is None else top,
            self.width - 2 * MARGIN if width is None else width,
        )

    def _advance(self, slide, top: float, height_emu: int) -> None:
        self._cursor[slide.slide_id] = max(
            self._cursor.get(slide.slide_id, MARGIN), top + height_emu / 914400 + GAP
        )

    def text(
        self,
        slide,
        text: str,
        *,
        size: float | None = None,
        bold: bool = False,
        left: float | None = None,
        top: float | None = None,
        width: float | None = None,
    ):
        """A text box; each line of `text` is a paragraph. Positions in inches."""
        left, top, width = self._place(slide, left, top, width)
        size = size or self.config.deck.size
        lines = text.split("\n")
        per_line = max(1, int(width * 72 / (size * CHAR_EM)))
        rows = sum(max(1, -(-len(line) // per_line)) for line in lines)
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Pt(size * 1.3 * rows))
        frame = box.text_frame
        frame.word_wrap = True
        for i, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            run = paragraph.add_run()
            run.text = line
            run.font.name, run.font.size, run.font.bold = self.config.deck.font, Pt(size), bold
        self._advance(slide, top, box.height)
        return box

    def math(
        self,
        slide,
        latex: str,
        *,
        route: str | None = None,
        size: float | None = None,
        left: float | None = None,
        top: float | None = None,
        width: float | None = None,
    ):
        """An equation by `route`, else the configured one. Positions in inches."""
        left, top, width = self._place(slide, left, top, width)
        settings = self.config.math
        size = size or settings.size
        route = route or settings.route
        font = settings.font if route == "omml" else settings.tex_font
        shape = add_equation(
            slide,
            latex,
            route=route,
            font=font,
            size_pt=size,
            left=Inches(left),
            top=Inches(top),
            width=Inches(width),
        )
        self._advance(slide, top, shape.height)
        return shape

    def figure(
        self,
        slide,
        path: str | Path,
        *,
        fallback: str | Path | None = None,
        left: float | None = None,
        top: float | None = None,
        width: float | None = None,
        height: float | None = None,
    ):
        """A picture; relative paths are taken from the configuration's folder. Positions in inches."""
        left, top, full = self._place(slide, left, top, width)
        root = self.config.root
        picture = add_figure(
            slide,
            root / path,
            left=Inches(left),
            top=Inches(top),
            width=Inches(width) if width else None,
            height=Inches(height) if height else None,
            fallback=root / fallback if fallback else None,
        )
        if width is None and height is None:
            fit(picture, Inches(full))
        self._advance(slide, top, picture.height)
        return picture

    def add(self, spec: dict) -> object:
        """A slide from one [[slide]] table."""
        if extra := set(spec) - SLIDE_KEYS:
            raise ValueError(f"unknown keys in [[slide]]: {', '.join(sorted(extra))}")
        slide = self.slide(spec.get("title", ""))
        for item in spec.get("item", []):
            kinds = [k for k in ITEM_KEYS if k in item]
            if len(kinds) != 1:
                raise ValueError(f"slide item needs one of text, math or figure: {item}")
            if extra := set(item) - {kinds[0]} - ITEM_KEYS[kinds[0]] - PLACE_KEYS:
                raise ValueError(f"unknown keys in {kinds[0]} item: {', '.join(sorted(extra))}")
            where = {k: item[k] for k in ("left", "top", "width") if k in item}
            if "text" in item:
                self.text(slide, item["text"], size=item.get("size"), bold=item.get("bold", False), **where)
            elif "math" in item:
                self.math(slide, item["math"], route=item.get("route"), size=item.get("size"), **where)
            else:
                self.figure(
                    slide, item["figure"], fallback=item.get("fallback"), height=item.get("height"), **where
                )
        return slide

    def save(self, path: str | Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)
        return path


def build(config: Config, out: str | Path) -> Path:
    """The deck described by `config`, saved to `out`."""
    deck = Deck(config)
    for spec in config.slides:
        deck.add(spec)
    return deck.save(out)
