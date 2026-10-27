"""LaTeX to Office math (OMML) that PowerPoint shows as an editable equation.

LaTeX -> MathML (latex2mathml) -> OMML (Office's MML2OMML.XSL). The stylesheet ships with
Microsoft Office and is not redistributable, so it is looked up in the Office folder or taken
from the MML2OMML environment variable. MathML is adjusted before the stylesheet runs and the
OMML after it, where the stylesheet's output sets differently in PowerPoint than in TeX.
"""

from __future__ import annotations

import copy
import os
from functools import cache
from pathlib import Path

from latex2mathml.converter import convert
from lxml import etree
from pptx.util import Pt

from .xml import NS, q

MATHML = "http://www.w3.org/1998/Math/MathML"
FUNCTIONS = {
    "sin", "cos", "tan", "cot", "sec", "csc", "exp", "log", "ln", "lg", "sinh", "cosh", "tanh",
    "coth", "arcsin", "arccos", "arctan", "arg", "det", "dim", "ker", "deg", "gcd", "max", "min",
    "sup", "inf", "lim",
}  # fmt: skip
NARY = set("∑∏∐∫∬∭∮∯∰⋃⋂⋁⋀⨁⨂⨀")
OPERAND_ENDS = set("=+−-±∓∝<>≤≥≈≡∼≠≪≫→⇒,;")
SCRIPTS = {f"{{{MATHML}}}{tag}" for tag in ("msub", "msup", "msubsup", "munder", "mover", "munderover")}
UPRIGHT = set("ΓΔΘΛΞΠΣΥΦΨΩ∇")
# Unicode spaces by width in em: thin (1/6 TeX \,), four-per-em, three-per-em, en.
SPACES = {" ": 0.167, " ": 0.25, " ": 0.333, " ": 0.5}
OFFICE_ROOTS = ("Program Files", "Program Files (x86)")


def stylesheet() -> Path | None:
    """MML2OMML.XSL from the MML2OMML variable or the Office installation, else None."""
    if env := os.environ.get("MML2OMML"):
        return Path(env)
    drive = Path(os.environ.get("SystemDrive", "C:") + os.sep)
    for root in OFFICE_ROOTS:
        for found in (drive / root / "Microsoft Office").glob("**/MML2OMML.XSL"):
            return found
    return None


@cache
def _transform() -> etree.XSLT:
    path = stylesheet()
    if path is None or not path.exists():
        raise FileNotFoundError("MML2OMML.XSL not found: install Microsoft Office or set MML2OMML")
    return etree.XSLT(etree.parse(str(path)))


def _m(tag: str) -> str:
    return f"{{{MATHML}}}{tag}"


def mathml(latex: str) -> etree._Element:
    """LaTeX -> display MathML, with the adjustments the stylesheet needs."""
    root = etree.fromstring(convert(latex, display="block").encode())
    apply_functions(root)
    group_nary_operands(root)
    upright_symbols(root)
    keep_spaces(root)
    return root


def upright_symbols(root: etree._Element) -> None:
    """PowerPoint sets every single-letter run in math italic, also capital Greek and nabla,
    which TeX sets upright: mark them as normal, which the stylesheet writes as plain style."""
    for node in list(root.iter(_m("mi"), _m("mo"))):
        if node.text and len(node.text) == 1 and node.text in UPRIGHT and node.get("mathvariant") is None:
            node.tag = _m("mi")
            node.set("mathvariant", "normal")


def keep_spaces(root: etree._Element) -> None:
    """The stylesheet drops <mspace>, so \\, and \\qquad vanish. Replace each by the Unicode
    space nearest its width, which PowerPoint keeps."""
    for space in list(root.iter(_m("mspace"))):
        width = space.get("width", "")
        em = float(width.removesuffix("em")) if width.endswith("em") else 0.0
        if em <= 0:
            space.getparent().remove(space)
            continue
        text = etree.Element(_m("mtext"))
        whole, rest = divmod(em, 1.0)
        text.text = " " * int(whole) + (
            min(SPACES, key=lambda s: abs(SPACES[s] - rest)) if rest > 0.05 else ""
        )
        space.getparent().replace(space, text)


def apply_functions(root: etree._Element) -> None:
    """latex2mathml writes \\sin as a bare <mi>sin</mi>, which the stylesheet turns into plain
    text, so PowerPoint sets "sinx". The stylesheet builds an m:func only from a row that opens
    with the name and an ApplyFunction operator, so the name, U+2061 and the argument get a row
    of their own."""
    for name in list(root.iter(_m("mi"))):
        if name.text not in FUNCTIONS:
            continue
        after = name.getnext()
        if after is not None and after.tag == _m("mo") and after.text == "⁡":
            after = after.getnext()
        if after is None:
            continue
        group = etree.Element(_m("mrow"))
        name.addprevious(group)
        apply = etree.SubElement(group, _m("mo"))
        apply.text = "⁡"
        group.insert(0, name)
        group.append(after)


def _fence(node: etree._Element) -> int:
    """+1 for an opening bracket, -1 for a closing one, also as the base of a script like ")^2"."""
    base = node[0] if node.tag in SCRIPTS and len(node) else node
    if base.tag != _m("mo") or not base.text or len(base.text) != 1:
        return 0
    return (base.text in "([{") - (base.text in ")]}")


def group_nary_operands(root: etree._Element) -> None:
    """The stylesheet takes only the next sibling of a sum or integral as its operand, so in
    \\sum_n J_n(S) the operand is J_n and "(S)" follows the closed n-ary, where PowerPoint sets
    a gap: "J_n (S)". The operand becomes everything up to the next relation or +/- outside
    brackets, as TeX reads it."""
    for node in list(root.iter(_m("mo"))):
        if node.text not in NARY:
            continue
        parent = node.getparent()
        op = parent if parent.tag in SCRIPTS and parent[0] is node else node
        operand, depth, after = [], 0, op.getnext()
        while after is not None:
            depth += _fence(after)
            if depth < 0 or (depth == 0 and after.tag == _m("mo") and after.text in OPERAND_ENDS):
                break
            operand.append(after)
            after = after.getnext()
        if len(operand) > 1:
            group = etree.Element(_m("mrow"))
            op.addnext(group)
            group.extend(operand)


def _extent(node: etree._Element) -> tuple[float, float]:
    """(height above, depth below) the baseline in em, roughly as Office math lays it out.
    PowerPoint sizes the box only when it renders, so the layout needs this guess."""
    tag = etree.QName(node).localname
    kids = [_extent(child) for child in node]
    if tag == "mfrac" and len(kids) == 2:
        (na, nd), (da, dd) = kids
        return 0.35 + na + nd, da + dd - 0.15
    if tag in ("msub", "msup", "msubsup", "munder", "mover", "munderover") and kids:
        big = node[0].text in NARY if len(node) and node[0].text else False
        a, d = (1.0, 0.45) if big else kids[0]
        under = "munder" in tag
        scripts = [sum(k) * 0.72 for k in kids[1:]]
        lower = scripts[0] if tag in ("msub", "msubsup", "munder", "munderover") and scripts else 0.0
        upper = scripts[-1] if tag in ("msup", "msubsup", "mover", "munderover") and scripts else 0.0
        if tag == "mover" and not big:
            return a + 0.15, d
        if big and (under or tag == "msubsup"):
            return a + upper + (0.1 if upper else 0), d + lower + (0.1 if lower else 0)
        return max(a, 0.35 + upper), max(d, 0.1 + lower)
    if tag in ("msqrt", "mroot") and kids:
        return max(k[0] for k in kids) + 0.15, max(k[1] for k in kids)
    if tag == "mtable":
        return sum(sum(k) for k in kids), 0.0
    if not kids:
        return 0.75, 0.25
    return max(k[0] for k in kids), max(k[1] for k in kids)


def height_pt(latex: str, size_pt: float) -> float:
    """The rendered height of `latex` at `size_pt`, an estimate."""
    above, below = _extent(mathml(latex))
    return (above + below) * size_pt * 1.15 + 7.2


def _bold(ch: str) -> str | None:
    """The plain letter behind a mathematical bold letter (U+1D400-1D433), else None."""
    code = ord(ch) - 0x1D400
    if 0 <= code < 26:
        return chr(ord("A") + code)
    if 26 <= code < 52:
        return chr(ord("a") + code - 26)
    return None


def split_bold(math: etree._Element) -> None:
    """PowerPoint drops mathematical bold letters inside a merged run such as "I(𝐐)": write them
    as plain letters in runs of their own with the bold math style, as PowerPoint stores bold."""
    for run in list(math.iter(q("m:r"))):
        text = run.find(q("m:t"))
        if text is None or not text.text or not any(_bold(c) for c in text.text):
            continue
        pieces: list[tuple[str, bool]] = []
        for ch in text.text:
            plain = _bold(ch)
            bold = plain is not None
            plain = plain or ch
            if pieces and pieces[-1][1] == bold:
                pieces[-1] = (pieces[-1][0] + plain, bold)
            else:
                pieces.append((plain, bold))
        for piece, bold in pieces:
            new = copy.deepcopy(run)
            new.find(q("m:t")).text = piece
            if bold:
                rpr = new.find(q("m:rPr"))
                if rpr is None:
                    rpr = etree.Element(q("m:rPr"))
                    new.insert(0, rpr)
                etree.SubElement(rpr, q("m:sty"), {q("m:val"): "b"})
            run.addprevious(new)
        run.getparent().remove(run)


def omml(latex: str, *, font: str = "Cambria Math", size_pt: float = 24) -> etree._Element:
    """LaTeX -> an m:oMath element whose runs carry PowerPoint run properties."""
    math = _transform()(mathml(latex)).getroot()
    if math.tag == q("m:oMathPara"):
        math = math.find(q("m:oMath"))
    for rpr in list(math.iter(q("w:rPr"))):
        rpr.getparent().remove(rpr)
    split_bold(math)
    for run in math.iter(q("m:r")):
        rpr = etree.Element(q("a:rPr"), lang="en-US", sz=str(round(size_pt * 100)))
        etree.SubElement(rpr, q("a:latin"), typeface=font)
        etree.SubElement(rpr, q("a:cs"), typeface=font)
        text = run.find(q("m:t"))
        (text.addprevious if text is not None else run.append)(rpr)
    return math


def add_equation(slide, latex: str, *, font: str, size_pt: float, left, top, width):
    """A text box holding `latex` as one left-aligned display equation; the LaTeX source is the
    fallback text for viewers without Office math."""
    shape = slide.shapes.add_textbox(left, top, width, Pt(height_pt(latex, size_pt)))
    frame = shape.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]._p
    alt = etree.SubElement(paragraph, q("mc:AlternateContent"), nsmap={"mc": NS["mc"]})
    choice = etree.SubElement(alt, q("mc:Choice"), Requires="a14", nsmap={"a14": NS["a14"]})
    para = etree.SubElement(etree.SubElement(choice, q("a14:m")), q("m:oMathPara"), nsmap={"m": NS["m"]})
    etree.SubElement(etree.SubElement(para, q("m:oMathParaPr")), q("m:jc"), {q("m:val"): "left"})
    para.append(omml(latex, font=font, size_pt=size_pt))
    fallback = etree.SubElement(etree.SubElement(alt, q("mc:Fallback")), q("a:r"))
    etree.SubElement(fallback, q("a:rPr"), lang="en-US", sz=str(round(size_pt * 100)))
    etree.SubElement(fallback, q("a:t")).text = latex
    end = paragraph.find(q("a:endParaRPr"))
    if end is not None:
        paragraph.remove(end)
        paragraph.append(end)
    return shape
